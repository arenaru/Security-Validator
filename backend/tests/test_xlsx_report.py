"""
Tests for the XLSX report builder.

These pin the per-module column headers so that a drift between
``backend/utils/report_columns.py`` and ``getColumnsForModule`` in
``frontend/src/components/ResultsTable.tsx`` surfaces as a failing test rather
than a silently wrong report.
"""
from io import BytesIO

from openpyxl import load_workbook

from backend.models.scan_models import ModuleResult, ResultStatus, ScanJob, ScanOptions
from backend.services.scan_service import InMemoryScanStore, ScanService


RESPONSE_CODE = "Response Code Check"

# The six columns the UI renders for Response Code Check, in order.
EXPECTED_RESPONSE_CODE_HEADERS = [
    "No",
    "URL",
    "Status Code",
    "Final Code",
    "Final URL",
    "Message",
]

GENERIC_HEADERS = [
    "module",
    "target",
    "status",
    "details",
    "severity",
    "code",
    "vuln_name",
]


def _response_code_result(url, status_code, final_code, final_url, message):
    return ModuleResult(
        module=RESPONSE_CODE,
        target=url,
        status=ResultStatus.INFO,
        details=message,
        raw={
            "URL": url,
            "Status Code": status_code,
            "Final Code": final_code,
            "Final URL": final_url,
            "Reason": "",
            "Redirects": 1 if status_code != final_code else 0,
            "Message": message,
        },
    )


def _build_workbook(modules, results):
    """Persist a finished job directly and render it, skipping the live scan."""
    store = InMemoryScanStore()
    job = ScanJob(
        targets=["example.com"],
        modules=modules,
        options=ScanOptions(),
    )
    job.results = results
    store.save(job)

    service = ScanService(store=store)
    payload = service.build_xlsx_report(job.scan_id)
    return load_workbook(BytesIO(payload))


def _headers(sheet):
    return [cell.value for cell in sheet[1]]


def test_response_code_sheet_matches_ui_columns():
    workbook = _build_workbook(
        [RESPONSE_CODE],
        {
            RESPONSE_CODE: [
                _response_code_result(
                    "https://ok.example.com", 200, 200,
                    "https://ok.example.com", "HTTP 200 OK",
                ),
            ]
        },
    )

    assert _headers(workbook[RESPONSE_CODE]) == EXPECTED_RESPONSE_CODE_HEADERS


def test_redirect_row_keeps_both_first_and_final_code():
    workbook = _build_workbook(
        [RESPONSE_CODE],
        {
            RESPONSE_CODE: [
                _response_code_result(
                    "https://unit.example.com", 308, 200,
                    "https://unit.example.com/product", "HTTP 308 -> 200 OK",
                ),
            ]
        },
    )

    sheet = workbook[RESPONSE_CODE]
    row = {header: sheet.cell(row=2, column=index + 1).value
           for index, header in enumerate(_headers(sheet))}

    assert row["No"] == 1
    assert row["URL"] == "https://unit.example.com"
    assert row["Status Code"] == 308
    assert row["Final Code"] == 200
    assert row["Final URL"] == "https://unit.example.com/product"


def test_row_numbers_are_sequential():
    workbook = _build_workbook(
        [RESPONSE_CODE],
        {
            RESPONSE_CODE: [
                _response_code_result(
                    f"https://host{n}.example.com", 200, 200,
                    f"https://host{n}.example.com", "HTTP 200 OK",
                )
                for n in range(3)
            ]
        },
    )

    sheet = workbook[RESPONSE_CODE]
    assert [sheet.cell(row=r, column=1).value for r in (2, 3, 4)] == [1, 2, 3]


def test_http_code_cells_are_tinted_by_bucket():
    """
    Emerald for codes that answered, amber for 4xx/5xx, slate for no code.

    N/A must not be tinted green: the original UI rule tested only
    startswith('4'/'5'), so an unreachable host rendered as a success.
    """
    workbook = _build_workbook(
        [RESPONSE_CODE],
        {
            RESPONSE_CODE: [
                _response_code_result(
                    "https://ok.example.com", 200, 200,
                    "https://ok.example.com", "HTTP 200 OK",
                ),
                _response_code_result(
                    "https://missing.example.com", 404, 404,
                    "https://missing.example.com", "HTTP 404 Not Found",
                ),
                _response_code_result(
                    "https://down.example.com", "N/A", "N/A",
                    "https://down.example.com", "Connection refused",
                ),
            ]
        },
    )

    sheet = workbook[RESPONSE_CODE]
    status_code_column = _headers(sheet).index("Status Code") + 1

    fills = [sheet.cell(row=r, column=status_code_column).fill.start_color.rgb
             for r in (2, 3, 4)]

    ok_fill, error_fill, none_fill = fills
    assert ok_fill == "FFD1FAE5"
    assert error_fill == "FFFEF3C7"
    assert none_fill == "FFE2E8F0"
    assert len({ok_fill, error_fill, none_fill}) == 3


def test_final_code_column_is_tinted_independently():
    """
    The two code columns are tinted from their own values, not from one shared
    verdict: a 503 -> 200 row shows amber on the first hop and emerald on the
    final code.
    """
    workbook = _build_workbook(
        [RESPONSE_CODE],
        {
            RESPONSE_CODE: [
                _response_code_result(
                    "https://unit.example.com", 503, 200,
                    "https://unit.example.com/product", "HTTP 503 -> 200",
                ),
            ]
        },
    )

    sheet = workbook[RESPONSE_CODE]
    headers = _headers(sheet)
    status_cell = sheet.cell(row=2, column=headers.index("Status Code") + 1)
    final_cell = sheet.cell(row=2, column=headers.index("Final Code") + 1)

    assert status_cell.fill.start_color.rgb == "FFFEF3C7"
    assert final_cell.fill.start_color.rgb == "FFD1FAE5"


def test_headers_are_written_even_when_module_has_no_rows():
    """
    A module that timed out stores an empty list. The sheet must still carry
    its headers; pd.DataFrame([]) would otherwise emit a blank sheet.
    """
    workbook = _build_workbook([RESPONSE_CODE], {RESPONSE_CODE: []})

    sheet = workbook[RESPONSE_CODE]
    assert _headers(sheet) == EXPECTED_RESPONSE_CODE_HEADERS
    assert sheet.max_row == 1


def test_module_without_a_spec_falls_back_to_generic_columns():
    workbook = _build_workbook(
        ["PHP Version Disclosure"],
        {
            "PHP Version Disclosure": [
                ModuleResult(
                    module="PHP Version Disclosure",
                    target="https://example.com",
                    status=ResultStatus.WARNING,
                    details="PHP 7.4 disclosed",
                    raw={"URL": "https://example.com"},
                )
            ]
        },
    )

    sheet = workbook["PHP Version Disclosure"]
    assert _headers(sheet) == GENERIC_HEADERS
    assert sheet.cell(row=2, column=2).value == "https://example.com"
    assert sheet.cell(row=2, column=3).value == "warning"


def test_missing_payload_key_renders_as_na():
    """A module row lacking a mapped key must not produce an empty cell."""
    workbook = _build_workbook(
        [RESPONSE_CODE],
        {
            RESPONSE_CODE: [
                ModuleResult(
                    module=RESPONSE_CODE,
                    target="https://partial.example.com",
                    status=ResultStatus.INFO,
                    details="-",
                    raw={"URL": "https://partial.example.com"},
                )
            ]
        },
    )

    sheet = workbook[RESPONSE_CODE]
    headers = _headers(sheet)
    assert sheet.cell(row=2, column=headers.index("Final Code") + 1).value == "N/A"


def test_summary_sheet_lists_every_requested_module():
    workbook = _build_workbook(
        [RESPONSE_CODE, "PHP Version Disclosure"],
        {RESPONSE_CODE: [], "PHP Version Disclosure": []},
    )

    assert "summary" in workbook.sheetnames
    summary = workbook["summary"]
    modules = [summary.cell(row=r, column=1).value for r in (2, 3)]
    assert modules == [RESPONSE_CODE, "PHP Version Disclosure"]
